#!/usr/bin/env python3
"""
Build the ZIP -> Congress data the "Where does your representative stand on the
vaccine executive order?" section reads, from the source spreadsheet and the
House portrait archive.

    python3 scripts/build-congress-data.py \
        --xlsx      ~/Downloads/zip_to_congressional_district_119th_with_eo_stance.xlsx \
        --portraits ~/Downloads/house_member_portraits_119th

Writes, relative to the repo root:

    data/house-members.json   every House seat keyed by district code, plus the
                              FIPS-derived state-abbreviation -> map-name table
    data/zips/NN.json         the ZIP lookup, sharded by the ZIP's first two
                              digits so a lookup fetches ~20KB instead of ~2MB
    images/house/<file>.jpg   portraits, downscaled to 240px wide

Shard schema, kept terse because it ships to the browser:

    { "<zip>": [ "<state abbr>", [ [ "<district code>", <1 primary|0>, <pct|null> ], ... ] ] }

The district list is pre-sorted: primary district first, then by share of the
ZIP's land area descending, with blank shares last. The page renders it in the
order it is given.

Portraits are matched to seats by the district-code prefix of the filename
(TX-37_Doggett_Lloyd.jpg -> TX-37), never by member name -- punctuation,
accents, suffixes and compound surnames make name matching unreliable.

Structural problems (missing columns, an unmatched portrait, a seat whose
portrait is missing) abort the build rather than writing partial data. The
published baselines below are checked and reported but do not abort, so a
refreshed spreadsheet can legitimately move them.
"""

import argparse
import collections
import json
import os
import re
import shutil
import sys
from datetime import date

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

REQUIRED_COLUMNS = [
    "zip", "zip_state", "district_code", "state_fips", "is_primary_district",
    "pct_of_zip_land_area", "seat_status", "member_name", "member_party",
    "member_type", "bioguide_id", "member_website",
    "childhood_vaccine_eo_position",
]

# What the supplied 2026-09-12 workbook contains. Reported, not enforced.
BASELINE = {
    "rows": 47925,
    "zips": 41637,
    "districts": 441,
    "vacant": ["FL-20", "TX-23"],
    "stance": {"Opposes": 70, "Supports": 2, "No public stance identified": 367},
}

PORTRAIT_PREFIX_RE = re.compile(r"^([A-Z]{2}-(?:\d{2}|AL))_")

# The three values the spreadsheet uses. Anything else is a data error.
VALID_POSITIONS = {"Opposes", "Supports", "No public stance identified"}


def die(msg):
    sys.exit("build-congress-data: " + msg)


def read_rows(xlsx_path):
    try:
        import openpyxl
    except ImportError:
        die("openpyxl is required to read the workbook: pip3 install openpyxl")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else "" for h in next(it)]
    except StopIteration:
        die("workbook sheet %r is empty" % wb.sheetnames[0])

    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        die("workbook is missing required column(s): " + ", ".join(missing))

    rows = []
    for raw in it:
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        rows.append(dict(zip(header, raw)))
    if not rows:
        die("workbook has a header but no data rows")
    return rows


def clean(v):
    """Spreadsheet cell -> trimmed string, or None when blank."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def to_pct(v):
    s = clean(v)
    if s is None:
        return None
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def build_portrait_index(portraits_dir):
    """district code -> portrait filename, matched on the filename's prefix."""
    if not os.path.isdir(portraits_dir):
        die("portrait directory not found: " + portraits_dir)

    index, unparsed = {}, []
    for name in sorted(os.listdir(portraits_dir)):
        if not name.lower().endswith((".jpg", ".jpeg")):
            continue
        m = PORTRAIT_PREFIX_RE.match(name)
        if not m:
            unparsed.append(name)
            continue
        code = m.group(1)
        if code in index:
            die("two portraits claim district %s: %s and %s" % (code, index[code], name))
        index[code] = name

    if unparsed:
        die("portrait filename(s) do not start with a district code: " + ", ".join(unparsed))
    if not index:
        die("no portraits found in " + portraits_dir)
    return index


def copy_portraits(portraits_dir, index, out_dir, width=240, quality=82):
    """Downscale each portrait into the repo. Cards render these ~92px wide."""
    try:
        from PIL import Image
    except ImportError:
        die("Pillow is required to resize portraits: pip3 install Pillow")

    if os.path.isdir(out_dir):
        shutil.rmtree(out_dir)
    os.makedirs(out_dir)

    total = 0
    for code, name in sorted(index.items()):
        src = os.path.join(portraits_dir, name)
        dst = os.path.join(out_dir, name)
        with Image.open(src) as im:
            im = im.convert("RGB")
            if im.width > width:
                im = im.resize((width, round(im.height * width / im.width)), Image.LANCZOS)
            im.save(dst, "JPEG", quality=quality, optimize=True, progressive=True)
        total += os.path.getsize(dst)
    return total


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, help="the ZIP-to-district workbook")
    ap.add_argument("--portraits", required=True, help="the unzipped portrait directory")
    ap.add_argument("--skip-portraits", action="store_true",
                    help="rebuild the JSON only, leaving images/house alone")
    args = ap.parse_args()

    xlsx = os.path.expanduser(args.xlsx)
    portraits_dir = os.path.expanduser(args.portraits)
    if not os.path.isfile(xlsx):
        die("workbook not found: " + xlsx)

    print("reading %s" % xlsx)
    rows = read_rows(xlsx)
    print("  %s ZIP-district rows" % f"{len(rows):,}")

    portrait_index = build_portrait_index(portraits_dir)
    print("  %d portraits indexed by district code" % len(portrait_index))

    # ---- seats -------------------------------------------------------------
    seats, fips_to_abbr = {}, {}
    for r in rows:
        code = clean(r["district_code"])
        if not code:
            die("a row has no district_code")

        abbr = code.split("-")[0]
        fips = clean(r["state_fips"])
        if fips:
            fips_to_abbr[fips.zfill(2)] = abbr

        if code in seats:
            continue

        status = clean(r["seat_status"]) or "Unknown"
        vacant = status.lower() != "filled"
        position = clean(r["childhood_vaccine_eo_position"])
        if not vacant and position not in VALID_POSITIONS:
            die("seat %s has an unrecognised position %r (expected one of: %s)"
                % (code, position, ", ".join(sorted(VALID_POSITIONS))))

        seat = {
            "districtCode": code,
            "stateAbbr": abbr,
            "seatStatus": "Vacant" if vacant else "Filled",
        }
        if not vacant:
            portrait = portrait_index.get(code)
            if not portrait:
                die("filled seat %s has no portrait in %s" % (code, portraits_dir))
            seat.update({
                "name": clean(r["member_name"]),
                "party": clean(r["member_party"]),
                "memberType": clean(r["member_type"]) or "Representative",
                "bioguideId": clean(r["bioguide_id"]),
                "website": clean(r["member_website"]),
                "position": position,
                "portrait": "images/house/" + portrait,
            })
            if not seat["name"]:
                die("filled seat %s has no member_name" % code)
        seats[code] = seat

    orphans = sorted(set(portrait_index) - set(seats))
    if orphans:
        die("portrait(s) for district(s) not in the workbook: " + ", ".join(orphans))

    vacant = sorted(c for c, s in seats.items() if s["seatStatus"] == "Vacant")
    stance = collections.Counter(s["position"] for s in seats.values()
                                 if s["seatStatus"] == "Filled")
    filled = [s for s in seats.values() if s["seatStatus"] == "Filled"]
    bioguides = {s["bioguideId"] for s in filled if s["bioguideId"]}
    if len(bioguides) != len(filled):
        die("filled seats (%d) and distinct bioguide ids (%d) disagree -- a member "
            "is listed in two districts" % (len(filled), len(bioguides)))

    # ---- ZIP lookup --------------------------------------------------------
    by_zip = collections.defaultdict(list)
    for r in rows:
        z = clean(r["zip"])
        if not z:
            die("a row has no zip")
        z = z.zfill(5)
        if not z.isdigit() or len(z) != 5:
            die("row has a ZIP that is not five digits: %r" % z)
        by_zip[z].append(r)

    shards = collections.defaultdict(dict)
    cross_state = 0
    for z, group in by_zip.items():
        seen, entries = set(), []
        for r in group:
            code = clean(r["district_code"])
            if code in seen:          # deduplicate by district_code
                continue
            seen.add(code)
            entries.append((
                code,
                1 if (clean(r["is_primary_district"]) or "").upper() == "Y" else 0,
                to_pct(r["pct_of_zip_land_area"]),
            ))

        # primary first, then land-area share descending, blank shares last.
        entries.sort(key=lambda e: (-e[1], e[2] is None, -(e[2] or 0), e[0]))

        # zip_state is the ZIP's own state; 141 ZIPs touch a district across a
        # state line, so it is not always the primary district's prefix. The
        # senators and the map highlight follow zip_state.
        state = clean(group[0]["zip_state"]) or entries[0][0].split("-")[0]
        if any(e[0].split("-")[0] != state for e in entries):
            cross_state += 1

        shards[z[:2]][z] = [state, entries]

    # ---- write -------------------------------------------------------------
    generated = date.today().isoformat()

    state_names = {}
    topo_path = os.path.join(REPO, "data", "states-10m.json")
    with open(topo_path) as fh:
        topo = json.load(fh)
    for g in topo["objects"]["states"]["geometries"]:
        abbr = fips_to_abbr.get(str(g.get("id", "")).zfill(2))
        if abbr:
            state_names[abbr] = g["properties"]["name"]
    for abbr in sorted({s["stateAbbr"] for s in seats.values()}):
        state_names.setdefault(abbr, None)

    members_path = os.path.join(REPO, "data", "house-members.json")
    with open(members_path, "w") as fh:
        json.dump({
            "generatedAt": generated,
            "source": os.path.basename(xlsx),
            "note": ("119th Congress district boundaries -- who represents a ZIP today. "
                     "Not the 2026-election map. Built by scripts/build-congress-data.py."),
            "stateNames": state_names,
            "members": dict(sorted(seats.items())),
        }, fh, separators=(",", ":"), sort_keys=False)
        fh.write("\n")

    zips_dir = os.path.join(REPO, "data", "zips")
    if os.path.isdir(zips_dir):
        shutil.rmtree(zips_dir)
    os.makedirs(zips_dir)
    shard_bytes = 0
    for prefix, mapping in sorted(shards.items()):
        p = os.path.join(zips_dir, prefix + ".json")
        with open(p, "w") as fh:
            json.dump(dict(sorted(mapping.items())), fh, separators=(",", ":"))
            fh.write("\n")
        shard_bytes += os.path.getsize(p)

    portrait_bytes = 0
    if args.skip_portraits:
        print("  skipping portraits (--skip-portraits)")
    else:
        portrait_bytes = copy_portraits(
            portraits_dir, portrait_index, os.path.join(REPO, "images", "house"))

    # ---- report ------------------------------------------------------------
    def cmp(label, got, want):
        ok = "ok" if got == want else "CHANGED (workbook had %s)" % (want,)
        print("  %-34s %-28s %s" % (label, got, ok))

    print("\nwrote data/house-members.json  (%.0f KB)" % (os.path.getsize(members_path) / 1024))
    print("wrote data/zips/*.json         (%d shards, %.0f KB total, %.1f KB avg)"
          % (len(shards), shard_bytes / 1024, shard_bytes / 1024 / len(shards)))
    if portrait_bytes:
        print("wrote images/house/*.jpg       (%d files, %.1f MB)"
              % (len(portrait_index), portrait_bytes / 1048576))

    print("\nagainst the supplied workbook's baseline:")
    cmp("ZIP-district rows", len(rows), BASELINE["rows"])
    cmp("distinct ZIPs", len(by_zip), BASELINE["zips"])
    cmp("district codes", len(seats), BASELINE["districts"])
    cmp("vacant seats", ",".join(vacant) or "(none)", ",".join(BASELINE["vacant"]))
    for key in ("Opposes", "Supports", "No public stance identified"):
        cmp("House " + key, stance.get(key, 0), BASELINE["stance"][key])

    print("\nother facts worth knowing:")
    print("  %-34s %s" % ("filled seats (unique members)", len(filled)))
    print("  %-34s %s" % ("split ZIPs (>1 district)",
                          sum(1 for v in by_zip.values() if len({clean(r["district_code"]) for r in v}) > 1)))
    print("  %-34s %s" % ("ZIPs touching another state", cross_state))
    print("  %-34s %s" % ("most districts in one ZIP",
                          max(len(v[1]) for m in shards.values() for v in m.values())))
    unmapped = sorted(a for a, n in state_names.items() if not n)
    print("  %-34s %s" % ("no state in the map geometry", ",".join(unmapped) or "(none)"))


if __name__ == "__main__":
    main()
